import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def make_docx():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_title.add_run("PANDUAN INTEGRASI GOOGLE NOTEBOOKLM\n")
    r1.bold = True
    r1.font.size = Pt(18)
    r1.font.color.rgb = RGBColor(31, 78, 121)

    r2 = p_title.add_run("PEMBELAJARAN, UJIAN & PODCAST AUDIO AI DARI PROYEK ERPNEXT\n")
    r2.bold = True
    r2.font.size = Pt(13)
    r2.font.color.rgb = RGBColor(68, 114, 196)

    r3 = p_title.add_run("Studi Kasus: Meli Dwi Company (MDC) - Manufaktur Meja Belajar Kayu Minimalis\n\n")
    r3.italic = True
    r3.font.size = Pt(11)

    doc.add_heading("1. Mengenal Google NotebookLM", level=1)
    doc.add_paragraph(
        "Google NotebookLM (https://notebooklm.google.com) adalah platform AI canggih dari Google yang memungkinkan Anda "
        "mengunggah dokumen proyek ERP (laporan, bank soal, buku petunjuk, flowchart) untuk menghasilkan ringkasan otomatis, "
        "kartu tanya jawab ujian (study guide), serta rekaman podcast audio dua pembicara (Audio Overview) yang membahas tuntas "
        "seluruh transaksi manufaktur dan akuntansi perusahaan Anda."
    )

    doc.add_heading("2. Berkas Sumber yang Siap Diunggah", level=1)
    doc.add_paragraph(
        "Seluruh berkas berada di folder: c:\\Users\\melid\\OneDrive\\Documents\\SEVENTH SEMESTER\\ENTERPRISE RESOURCE PLANNING (ERP)\\ERP_Meli\\\n\n"
        "1. LAPORAN_AKHIR_HASIL_PRAKTIK.docx: Laporan resmi lengkap kinerja manufaktur, laba rugi Rp 1.007.500, dan neraca.\n"
        "2. SOAL_DAN_JAWABAN_JURNAL_ERP.docx: Kumpulan soal ujian, kasus perakitan meja, dan kunci jurnal debit/kredit.\n"
        "3. BUKU_PETUNJUK_PENGGUNAAN_ERPNEXT.docx: Buku panduan manual navigasi dan pengoperasian localhost.\n"
        "4. FLOWCHART_PROSES_MANUFAKTUR.docx: Diagram alir visual proses bisnis manufaktur.\n"
        "5. SKENARIO_PRAKTIK_MANUFAKTUR.docx: Skenario pelaksanaan modul per modul."
    )

    doc.add_heading("3. Tiga Langkah Cepat Integrasi ke NotebookLM", level=1)
    doc.add_paragraph(
        "Langkah 1: Buka situs https://notebooklm.google.com di Google Chrome dan masuk dengan akun Google Anda.\n"
        "Langkah 2: Klik '+ New Notebook' dan beri nama 'Praktik ERP Manufaktur - Meli Dwi Company'.\n"
        "Langkah 3: Klik 'Upload Files' pada panel kiri, lalu pilih berkas-berkas DOCX dari folder ERP_Meli di atas."
    )

    doc.add_heading("4. Fitur Podcast Audio AI (Audio Overview)", level=1)
    doc.add_paragraph(
        "Setelah dokumen diunggah, klik tombol 'Generate' pada bagian 'Audio Overview (Deep Dive conversation)' di panel kanan atas. "
        "Dalam 2–3 menit, AI akan menyusun podcast audio interaktif yang membahas seluruh proyek ERP Anda. Sangat efektif untuk didengarkan "
        "menjelang presentasi atau ujian mata kuliah ERP!"
    )

    out_path = "/home/frappe/frappe-bench/sites/PANDUAN_INTEGRASI_NOTEBOOKLM.docx"
    doc.save(out_path)
    print(f"[OK] {out_path} berhasil dibuat!", flush=True)

if __name__ == "__main__":
    make_docx()
